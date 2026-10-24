# -*- coding: utf-8 -*-
"""Kreuzberg 主路切换（升级计划拍板项：迁移到 Kreuzberg 统一解析框架）。

`NEUROVA_KB_PARSER=kreuzberg` 一键切主路：内容层归 kreuzberg，本管线保留
清洗/分块/标题编排；图片不走 kreuzberg（OCR 是引擎层责任，RapidOCR 编排不变）；
旧格式仍先 LibreOffice 转换；抽取失败诚实报错（可读原因），不静默回落双主路。
缺省 legacy——切换是环境配置动作，机制在此单源。
"""
from __future__ import annotations

import io

import pytest

from neurova.knowledge import document_pipeline as dp
from neurova.knowledge import parser_shadow


class _StubOcr:
    name = "stub-ocr"

    def available(self) -> bool:
        return True

    def recognize(self, image: bytes):
        return "OCR 文本", 0.9, self.name


def _docxBytes() -> bytes:
    import docx

    buf = io.BytesIO()
    document = docx.Document()
    document.add_heading("主路标题", level=1)
    document.add_paragraph("主路正文内容")
    document.save(buf)
    return buf.getvalue()


def _xlsxBytes() -> bytes:
    import openpyxl

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active.append(["指标", "数值"])
    wb.active.append(["营收", "1200万"])
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def kreuzbergMode(monkeypatch):
    monkeypatch.setenv(parser_shadow.PARSER_MODE_ENV, "kreuzberg")


class TestKreuzbergMainPath:
    def testDocxGoesThroughKreuzberg(self, kreuzbergMode):
        parsed = dp.parse_document(_docxBytes(), "main.docx")

        assert "主路正文内容" in parsed.content
        assert parsed.metadata["pipeline"] == "kreuzberg"
        assert parsed.pages  # 编排产物仍由本管线生成

    def testXlsxContentExtracted(self, kreuzbergMode):
        parsed = dp.parse_document(_xlsxBytes(), "table.xlsx")

        assert "营收" in parsed.content and "1200万" in parsed.content
        assert parsed.metadata["pipeline"] == "kreuzberg"

    def testImagesStayOnEngineOcr(self, kreuzbergMode, monkeypatch):
        """图片不经 kreuzberg：OCR 编排（引擎层）照旧。"""
        from PIL import Image

        monkeypatch.setattr(dp, "selectOcrEngine", lambda: _StubOcr())
        buf = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(buf, format="PNG")

        parsed = dp.parse_document(buf.getvalue(), "scan.png", ocr_mode="on")

        assert parsed.pages[0]["text"] == "OCR 文本"
        assert parsed.metadata["ocr_engine"] == "stub-ocr"
        assert parsed.metadata["pipeline"] == "structured"

    def testUnknownFormatFailsHonestly(self, kreuzbergMode):
        from neurova.knowledge.document_pipeline import DocumentParseError

        with pytest.raises(DocumentParseError) as excinfo:
            dp.parse_document(b"whatever", "mystery.xyz")

        assert "kreuzberg" in str(excinfo.value).lower() or "不认格式" in str(excinfo.value)

    def testEmptyExtractionFailsHonestly(self, kreuzbergMode, monkeypatch):
        from neurova.knowledge.document_pipeline import DocumentParseError

        monkeypatch.setattr(
            parser_shadow, "_loadExtractor",
            lambda: (_ for _ in ()).throw(ValueError("Kreuzberg 抽取结果为空")))

        with pytest.raises(DocumentParseError):
            dp.parse_document(_docxBytes(), "empty.docx")

    def testDefaultModeIsStillLegacy(self, monkeypatch):
        monkeypatch.delenv(parser_shadow.PARSER_MODE_ENV, raising=False)

        parsed = dp.parse_document(_docxBytes(), "plain.docx")

        assert parsed.metadata["pipeline"] == "structured"
        assert parser_shadow.parserMode() == "legacy"

    def testShadowModeStillAttachesReading(self, monkeypatch):
        monkeypatch.setenv(parser_shadow.PARSER_MODE_ENV, "kreuzberg_shadow")

        parsed = dp.parse_document(_docxBytes(), "shadow.docx")

        assert parsed.metadata["pipeline"] == "structured"
        assert parsed.metadata["parser_shadow"]["ok"] is True
