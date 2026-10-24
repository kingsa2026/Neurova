# -*- coding: utf-8 -*-
"""格式/框架补齐（升级计划 A7）：.xls 可选直读 + Kreuzberg 影子对比。

验收（计划 §5 Batch A7）：
- .xls 在 xlrd 可用时直读，产出与 xlsx 完全同形；缺席回落 LibreOffice，
  不可用时报可读的 legacy_office_unavailable——格式矩阵每格都有可解释失败；
- Kreuzberg 影子模式只挂对比读数（metadata["parser_shadow"]），主路返回
  逐字不变；影子缺依赖 ok=False 带原因；缺省模式不产生影子读数。
"""
from __future__ import annotations

import io
import sys
import types

import pytest

from neurova.knowledge import document_pipeline as dp
from neurova.knowledge import parser_shadow


def _docxBytes(text: str = "影子对比正文") -> bytes:
    import docx

    buf = io.BytesIO()
    document = docx.Document()
    document.add_heading("影子标题", level=1)
    document.add_paragraph(text)
    document.save(buf)
    return buf.getvalue()


@pytest.fixture()
def stubXlrd(monkeypatch):
    """桩 xlrd：BIFF 解析无需真库，判据是直读分支的形状与分流。"""
    sheet = types.SimpleNamespace(
        name="Sheet1", nrows=2, ncols=2,
        cell_value=lambda r, c: [["名次", "分值"], ["1", "99"]][r][c])
    book = types.SimpleNamespace(nsheets=1, sheet_by_index=lambda i: sheet)
    mod = types.ModuleType("xlrd")
    mod.open_workbook = lambda file_contents=None: book
    monkeypatch.setitem(sys.modules, "xlrd", mod)
    return mod


class TestXlsDirectRead:
    def testXlsReadsDirectlyWhenXlrdPresent(self, stubXlrd, monkeypatch):
        monkeypatch.setattr(dp, "_convertLegacyOffice", lambda *a, **k: (_ for _ in ()).throw(
            AssertionError("xlrd 在场时不允许走 LibreOffice")))

        parsed = dp.parse_document(b"biff-bytes", "table.xls")

        assert parsed.pages[0]["text"] == "[工作表: Sheet1]\n名次 | 分值\n1 | 99"
        assert parsed.metadata["page_count"] == 1
        assert parsed.source_type == "xls" or parsed.source_type == ""

    def testXlsFallsBackWithoutXlrd(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "xlrd", None)
        from neurova.attachment_parser import extract_xls_sheets

        assert extract_xls_sheets(b"x") is None
        called = {}

        def _fakeConvert(data, ext):
            # LibreOffice 语义：.xls → .xlsx（返回真 xlsx 字节走完整后续链）。
            import openpyxl

            called["hit"] = ext
            buf = io.BytesIO()
            wb = openpyxl.Workbook()
            wb.active.append(["转换后", "单元格"])
            wb.save(buf)
            return (buf.getvalue(), ".xlsx")

        monkeypatch.setattr(dp, "_convertLegacyOffice", _fakeConvert)

        parsed = dp.parse_document(b"biff-bytes", "table.xls")

        assert called.get("hit") == ".xls"
        assert "转换后" in parsed.content

    def testSheetsShapeMatchesXlsx(self, stubXlrd):
        from neurova.attachment_parser import extract_xls_sheets

        sheets = extract_xls_sheets(b"biff")

        assert sheets == [{"title": "Sheet1", "text": "名次 | 分值\n1 | 99"}]


class TestKreuzbergShadow:
    def testShadowModeAttachesReadingAndKeepsMainPath(self, monkeypatch):
        monkeypatch.setenv(parser_shadow.PARSER_MODE_ENV, "kreuzberg_shadow")

        parsed = dp.parse_document(_docxBytes(), "shadow.docx")

        reading = parsed.metadata.get("parser_shadow")
        assert isinstance(reading, dict)
        assert reading["engine"] == "kreuzberg"
        assert reading["ok"] is True
        assert reading["chars"] > 0
        assert reading["main_chars"] == len(parsed.content)
        assert parsed.content  # 主路照常产出
        assert parsed.metadata["page_count"] >= 1

    def testDefaultModeHasNoShadowReading(self, monkeypatch):
        monkeypatch.delenv(parser_shadow.PARSER_MODE_ENV, raising=False)

        parsed = dp.parse_document(_docxBytes(), "plain.docx")

        assert "parser_shadow" not in parsed.metadata

    def testShadowFailureIsRecordedNotFatal(self, monkeypatch):
        monkeypatch.setenv(parser_shadow.PARSER_MODE_ENV, "kreuzberg_shadow")
        monkeypatch.setattr(
            parser_shadow, "_loadExtractor",
            lambda: (_ for _ in ()).throw(RuntimeError("kreuzberg 内核不可用")))

        parsed = dp.parse_document(_docxBytes(), "shadow.docx")

        reading = parsed.metadata["parser_shadow"]
        assert reading["ok"] is False
        assert "kreuzberg" in reading["error"]

    def testShadowImportFailureKeepsMainPathAlive(self, monkeypatch):
        """Kreuzberg 包缺席：影子如实记原因，主路解析不受任何影响。"""
        monkeypatch.setenv(parser_shadow.PARSER_MODE_ENV, "kreuzberg_shadow")
        realImport = __builtins__.__import__ if hasattr(__builtins__, "__import__") else __import__

        def _blocked(name, *args, **kwargs):
            if name == "kreuzberg" or name.startswith("kreuzberg."):
                raise ModuleNotFoundError("No module named 'kreuzberg'")
            return realImport(name, *args, **kwargs)

        monkeypatch.setattr("builtins.__import__", _blocked)

        parsed = dp.parse_document(_docxBytes(), "shadow.docx")

        reading = parsed.metadata["parser_shadow"]
        assert reading["ok"] is False
        assert reading["error"]
        assert parsed.content
