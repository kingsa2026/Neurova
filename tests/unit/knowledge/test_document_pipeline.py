"""文档标准化流水线：清洗、页码/标题、OCR、旧格式与失败原因。"""
from __future__ import annotations

import io

import pytest

from neurova.knowledge import document_pipeline as dp


class FakeOcr:
    name = "fake-ocr"

    def __init__(self, text="OCR 文本", available=True):
        self.text = text
        self._available = available

    def available(self):
        return self._available

    def recognize(self, image):
        if not self._available:
            raise dp.DocumentParseError("ocr_unavailable")
        return self.text, 0.91, self.name


def _docx_bytes():
    import docx

    document = docx.Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("正文内容")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes():
    import openpyxl

    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "配置"
    first.append(["键", "值"])
    first.append(["port", "9527"])
    second = workbook.create_sheet("说明")
    second.append(["说明正文"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx_bytes():
    from pptx import Presentation

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = "第一页标题"
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "第二页标题"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_clean_pages_removes_repeated_header_footer_and_page_number():
    pages = [
        "公司机密\n第一页正文\n第 1 页",
        "公司机密\n第二页正文\n第 2 页",
    ]
    cleaned = dp._clean_pages(pages)
    assert cleaned == ["第一页正文", "第二页正文"]


def test_markdown_heading_and_page_metadata():
    parsed = dp.parse_document(
        "# 标题\n\n正文内容\n\n## 小节\n\n更多内容".encode("utf-8"),
        "note.md",
    )
    assert parsed.source_type == "md"
    assert parsed.content.startswith("# 标题")
    assert parsed.pages[0]["page"] == 1
    assert parsed.chunks
    assert parsed.chunks[0]["page"] == 1
    assert parsed.chunks[0]["heading_path"] == ["标题"]


def test_docx_preserves_heading_path():
    parsed = dp.parse_document(_docx_bytes(), "需求.docx")
    assert "第一章" in parsed.content
    assert parsed.metadata["page_count"] == 1
    assert any(chunk["heading_path"] == ["第一章"] for chunk in parsed.chunks)


def test_xlsx_sheets_become_pages_with_sheet_heading():
    parsed = dp.parse_document(_xlsx_bytes(), "配置.xlsx")
    assert parsed.metadata["page_count"] == 2
    assert "键 | 值" in parsed.pages[0]["text"]
    assert any(chunk["heading_path"] == ["配置"] for chunk in parsed.chunks)
    assert any(chunk["heading_path"] == ["说明"] for chunk in parsed.chunks)


def test_pptx_slides_become_pages():
    parsed = dp.parse_document(_pptx_bytes(), "演示.pptx")
    assert parsed.metadata["page_count"] == 2
    assert "第一页标题" in parsed.pages[0]["text"]
    assert any(chunk["heading_path"] == ["幻灯片 1"] for chunk in parsed.chunks)


def test_pdf_page_mapping(monkeypatch):
    monkeypatch.setattr(
        dp, "extract_pdf_pages", lambda data: ["第一页正文", "第二页正文"]
    )
    parsed = dp.parse_document(b"%PDF-1.4", "doc.pdf", ocr_mode="off")
    assert parsed.metadata["page_count"] == 2
    assert parsed.pages[1]["text"] == "第二页正文"
    assert any(chunk["page"] == 2 for chunk in parsed.chunks)


def test_image_uses_ocr_engine():
    parsed = dp.parse_document(
        b"fake-image-bytes", "scan.png", ocr_engine=FakeOcr("扫描件文字")
    )
    assert parsed.content == "扫描件文字"
    assert parsed.metadata["ocr_used"] is True
    assert parsed.metadata["ocr_engine"] == "fake-ocr"
    assert parsed.metadata["ocr_confidence_avg"] == 0.91


def test_image_without_ocr_is_honest():
    with pytest.raises(dp.DocumentParseError) as caught:
        dp.parse_document(b"image", "scan.png", ocr_engine=FakeOcr(available=False))
    assert caught.value.reason == "ocr_unavailable"


def test_legacy_office_without_binary_returns_reason(monkeypatch):
    monkeypatch.setattr(dp, "_officeBinary", lambda: "")
    with pytest.raises(dp.DocumentParseError) as caught:
        dp.parse_document(b"\xd0\xcf\x11\xe0", "legacy.doc")
    assert caught.value.reason == "legacy_office_unavailable"


def test_unsupported_format_is_honest():
    with pytest.raises(dp.DocumentParseError) as caught:
        dp.parse_document(b"PK\x05\x06", "archive.zip")
    assert caught.value.reason == "unsupported_format"


def test_legacy_engine_is_runtime_rollback_switch():
    parsed = dp.parse_document(
        "# 标题\n\n旧管线正文".encode("utf-8"), "note.md", engine_mode="legacy"
    )

    assert parsed.metadata["pipeline"] == "legacy"
    assert parsed.metadata["page_count"] == 1
    assert "旧管线正文" in parsed.content
    assert parsed.warnings == ["legacy_pipeline"]


def test_docx_core_properties_are_extracted():
    from docx import Document

    document = Document()
    document.add_paragraph("正文")
    document.core_properties.title = "核心标题"
    document.core_properties.author = "张三"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = dp.parse_document(buffer.getvalue(), "meta.docx")

    assert parsed.metadata["title"] == "核心标题"
    assert parsed.metadata["author"] == "张三"


def test_repository_reparse_keeps_page_chunks_and_revision(tmp_path):
    from neurova.knowledge.repository import KnowledgeRepository

    repo = KnowledgeRepository(str(tmp_path / "kb"))
    item = repo.create_knowledge(
        "agent", "旧文", "旧正文", owner_user_id="u1", tenant_id="t1",
        document_id="doc-keep", source_type="md",
        chunks=[{"index": 0, "page": 1, "content": "旧正文"}],
    )
    parsed = dp.parse_document(
        "# 新标题\n\n新正文".encode("utf-8"), "n.md", engine_mode="structured"
    )

    ok = repo.replaceDocumentParsing("agent", item["knowledge_id"], parsed)

    updated = repo.get_item("agent", item["knowledge_id"])
    assert ok is True
    assert updated["content"] != "旧正文"
    assert updated["document_id"] == "doc-keep"
    assert updated["page_count"] == 1
    assert updated["revisions"][-1]["reason"] == "document_reparse"
