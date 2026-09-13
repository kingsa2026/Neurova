"""
P0-1 file_parse 文档理解工具（TDD 先红后绿）。

缺口（docs/Neurova_Agent工具扩展计划_2026-09-12.md P0-1）：file_read 只按
文本解码，PDF/Office 二进制文档 agent 只能看文件名。attachment_parser.
extract_attachment_text 已有完整文档解析（docx/xlsx/pptx/pdf/csv/html/rtf/
odf）但仅被附件链路消费——本工具将其接入内置工具面。计划原定 markitdown
就地修正为复用已验证模块（零新依赖，"只提升不下降"约束）。
"""

from unittest.mock import Mock

import pytest


def _make_executor(tmp_path):
    from neurova.tool_executor import ToolExecutor

    agent = Mock()
    agent._skill_registry = Mock()
    agent.tool_router = Mock()
    agent.tool_memory = Mock()
    agent.tool_lifecycle = Mock()
    agent.skill_packer = Mock()
    agent.config = Mock()
    agent.memory_manager = Mock()
    agent.asr_manager = None
    agent.tts_manager = None
    exe = ToolExecutor(agent)
    exe._agent.workspace_path = tmp_path  # 相对路径锚定工作区（09-08 契约）
    return exe


# ═══════════════════════════════════════════════════════════════
# 注册不变量（schema ⊆ 分派表 ⊆ 真实方法）
# ═══════════════════════════════════════════════════════════════

class TestFileParseRegistration:
    def test_schema_registered(self):
        from neurova.builtin_tools import _BUILTIN_SCHEMAS

        assert "file_parse" in _BUILTIN_SCHEMAS
        props = _BUILTIN_SCHEMAS["file_parse"]["parameters"]["properties"]
        assert "file_path" in props

    def test_dispatch_has_executable(self):
        from neurova.tool_executor import ToolExecutor

        assert "file_parse" in ToolExecutor._builtin_dispatch
        method = getattr(ToolExecutor, ToolExecutor._builtin_dispatch["file_parse"], None)
        assert method is not None and callable(method)


# ═══════════════════════════════════════════════════════════════
# 真实解析（python-docx/openpyxl/python-pptx 均为既有依赖）
# ═══════════════════════════════════════════════════════════════

class TestFileParseRealDocuments:
    @pytest.mark.asyncio
    async def test_docx(self, tmp_path):
        import docx

        doc = docx.Document()
        doc.add_paragraph("项目验收报告")
        doc.add_paragraph("第三章 结论")
        f = tmp_path / "report.docx"
        doc.save(str(f))

        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "report.docx"}
        )
        assert "error" not in result, result
        assert "项目验收报告" in result["content"]
        assert "第三章 结论" in result["content"]
        assert result["format"] == "docx"

    @pytest.mark.asyncio
    async def test_xlsx(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "预算"
        ws.append(["科目", "金额"])
        ws.append(["服务器", 42000])
        f = tmp_path / "budget.xlsx"
        wb.save(str(f))

        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "budget.xlsx"}
        )
        assert "error" not in result, result
        assert "科目" in result["content"] and "42000" in result["content"]
        assert "[工作表: 预算]" in result["content"]

    @pytest.mark.asyncio
    async def test_pptx(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Neurova 路线图"
        f = tmp_path / "deck.pptx"
        prs.save(str(f))

        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "deck.pptx"}
        )
        assert "error" not in result, result
        assert "Neurova 路线图" in result["content"]
        assert "[幻灯片 1]" in result["content"]

    @pytest.mark.asyncio
    async def test_plain_text_via_file_parse(self, tmp_path):
        f = tmp_path / "notes.md"
        f.write_text("# 标题\n正文", encoding="utf-8")
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "notes.md"}
        )
        assert "error" not in result, result
        assert "# 标题" in result["content"]


# ═══════════════════════════════════════════════════════════════
# 诚实降级（禁止静默失败）
# ═══════════════════════════════════════════════════════════════

class TestFileParseHonestDegradation:
    @pytest.mark.asyncio
    async def test_unsupported_format_reports_reason(self, tmp_path):
        f = tmp_path / "old.ppt"  # 旧 OLE2 格式明确不支持（见模块注释）
        f.write_bytes(b"\xd0\xcf\x11\xe0junkjunk")
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "old.ppt"}
        )
        assert "error" in result
        assert "unsupported_format" in result["error"]

    @pytest.mark.asyncio
    async def test_corrupt_pdf_does_not_raise(self, tmp_path):
        f = tmp_path / "broken.pdf"
        f.write_bytes(b"%PDF-1.4 not a real pdf at all")
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "broken.pdf"}
        )
        # 诚实报错而非异常，且指明解析失败类别
        assert "error" in result and "pdf" in result["error"].lower()

    @pytest.mark.asyncio
    async def test_missing_file(self, tmp_path):
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "nope.docx"}
        )
        assert "error" in result

    @pytest.mark.asyncio
    async def test_relative_escape_guard(self, tmp_path):
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "../outside.pdf"}
        )
        assert "error" in result and "路径越界" in result["error"]

    @pytest.mark.asyncio
    async def test_size_guard(self, tmp_path):
        f = tmp_path / "huge.pdf"
        f.write_bytes(b"%PDF" + b"\0" * (51 * 1024 * 1024))
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "huge.pdf"}
        )
        assert "error" in result and "过大" in result["error"]


# ═══════════════════════════════════════════════════════════════
# 截断闸门
# ═══════════════════════════════════════════════════════════════

class TestFileParseTruncation:
    @pytest.mark.asyncio
    async def test_max_chars_truncates_with_flag(self, tmp_path):
        f = tmp_path / "big.txt"
        f.write_text("x" * 5000, encoding="utf-8")
        result = await _make_executor(tmp_path)._execute_builtin_tool(
            "file_parse", {"file_path": "big.txt", "max_chars": 100}
        )
        assert "error" not in result, result
        assert len(result["content"]) == 100
        assert result["truncated"] is True
        assert result["total_chars"] >= 5000
